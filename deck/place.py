# Put the two animations into a copy of the deck: a video band under the title, the cards as a compact row.
import copy, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from lxml import etree

SRC = sys.argv[1]
OUT = sys.argv[2]
prs = Presentation(SRC)

def card_groups(slide, start, per=8):
    shapes = list(slide.shapes)
    names = [s.name for s in shapes]
    i = names.index(start)
    out = []
    while i + per <= len(shapes) and shapes[i].shape_type == 1 and shapes[i].name.startswith("Shape") and shapes[i].width > Inches(2):
        out.append(shapes[i:i + per]); i += per
    return out

def size(shape, pt):
    for p in shape.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(pt)

def place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)

TIMING = '''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:par><p:cTn id="2" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="5" dur="24000" fill="hold"/><p:tgtEl><p:spTgt spid="SPID"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par><p:video><p:cMediaNode vol="80000" mute="1"><p:cTn id="6" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="SPID"/></p:tgtEl></p:cMediaNode></p:video></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''

def add_video(slide, mp4, poster, x, y, w, h):
    mv = slide.shapes.add_movie(mp4, Inches(x), Inches(y), Inches(w), Inches(h), poster_frame_image=poster, mime_type="video/mp4")
    mv.name = "Animation"
    sld = slide._element
    for old in sld.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}timing"): sld.remove(old)
    timing = etree.fromstring(TIMING.replace("SPID", str(mv.shape_id)))
    ext = sld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}extLst")
    if ext is not None: ext.addprevious(timing)
    else: sld.append(timing)
    return mv

# ---- slide 4: three cards, band 11.56 wide ----
s4 = prs.slides[3]
add_video(s4, "slide4-old-toolkit.mp4", "poster4.jpg", 0.90, 1.92, 11.56, 11.56 / 3)
for g in card_groups(s4, "Shape 2"):
    bg, ring, icon, title, desc, rule, label, value = g
    x0, w = bg.left / 914400, bg.width / 914400
    place(bg, x0, 5.90, w, 1.06)
    place(ring, x0 + 0.18, 6.03, 0.44, 0.44); place(icon, x0 + 0.18, 6.03, 0.44, 0.44); size(icon, 9 if len(icon.text_frame.text) < 3 else 8)
    place(title, x0 + 0.72, 5.97, w - 0.84, 0.30); size(title, 14)
    place(desc, x0 + 0.72, 6.25, w - 0.84, 0.26); size(desc, 10)
    place(rule, x0 + 0.18, 6.60, w - 0.36, 0)
    place(label, x0 + 0.18, 6.64, 0.78, 0.26); size(label, 8)
    place(value, x0 + 0.92, 6.64, w - 1.04, 0.26); size(value, 10)
    for t in (title, desc, label, value): t.text_frame.word_wrap = True

# ---- slide 7: four cards, band 11.72 wide ----
s7 = prs.slides[6]
add_video(s7, "slide7-new-toolkit.mp4", "poster7.jpg", 0.90, 1.90, 11.72, 11.72 / 3)
for g in card_groups(s7, "Shape 2"):
    bg, ring, icon, title, model, rule, label, value = g
    x0, w = bg.left / 914400, bg.width / 914400
    place(bg, x0, 5.88, w, 1.12)
    place(ring, x0 + 0.15, 5.99, 0.40, 0.40); place(icon, x0 + 0.15, 5.99, 0.40, 0.40); size(icon, 8 if len(icon.text_frame.text) < 4 else 7)
    place(title, x0 + 0.62, 5.93, w - 0.72, 0.28); size(title, 12.5)
    place(model, x0 + 0.62, 6.19, w - 0.72, 0.22); size(model, 8)
    place(rule, x0 + 0.15, 6.47, w - 0.30, 0)
    place(label, x0 + 0.15, 6.50, 1.0, 0.18); size(label, 7.5)
    place(value, x0 + 0.15, 6.63, w - 0.30, 0.32); size(value, 9)
    for t in (title, model, label, value): t.text_frame.word_wrap = True

prs.save(OUT)
print("saved", OUT)
